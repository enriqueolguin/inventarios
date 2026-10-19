# -*- coding: utf-8 -*-
"""
Genera 3 libros de Excel (.xlsx) con propuestas de estructura de tablas para
inventariar servidores y VMs, a partir de documento/estructura.txt:

  Propuesta A - Resumen general + fichas por tema (una fila por servidor/VM)
  Propuesta B - Una hoja por cada sección del txt (13 hojas, orden original)
  Propuesta C - Hojas por tipo de equipo (físicos / VMs) + temas clave

Encabezados en negrita con color sutil, bordes en todas las celdas de trabajo,
autofiltro y primera fila congelada.
"""

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RUTA_SALIDA_A = "documento/Memoria_Servidores_Propuesta_A_Resumen_y_Fichas.xlsx"
RUTA_SALIDA_B = "documento/Memoria_Servidores_Propuesta_B_Una_hoja_por_seccion.xlsx"
RUTA_SALIDA_C = "documento/Memoria_Servidores_Propuesta_C_Por_tipo_de_equipo.xlsx"

FUENTE_NEGrita = Font(bold=True)

# Formato: bordes finos en todas las celdas y color sutil en encabezados
FILAS_PREFORMATEADAS = 100  # filas vacías con borde, listas para rellenar
RELLENO_ENCABEZADO = PatternFill(fill_type="solid", fgColor="DDEBF7")  # azul claro sutil
BORDE_DELGADO = Side(style="thin")
BORDE = Border(left=BORDE_DELGADO, right=BORDE_DELGADO, top=BORDE_DELGADO, bottom=BORDE_DELGADO)


# ---------------------------------------------------------------- utilidades
def crear_hoja(wb, nombre, columnas, primera=False):
    """Crea una hoja con fila de encabezados en negrita, autofiltro y ancho ajustado."""
    ws = wb.active if primera else wb.create_sheet()
    ws.title = nombre
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = FUENTE_NEGrita
        celda.fill = RELLENO_ENCABEZADO
        celda.border = BORDE
    # bordes en las filas vacías siguientes, listas para rellenar
    for fila in range(2, 2 + FILAS_PREFORMATEADAS):
        for col in range(1, len(columnas) + 1):
            ws.cell(row=fila, column=col).border = BORDE
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{1 + FILAS_PREFORMATEADAS}"
    ws.freeze_panes = "A2"
    for i, col in enumerate(columnas, start=1):
        ancho = min(max(len(col) + 4, 14), 42)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    return ws


# ------------------------------------------------------------------ campos
# Campos tomados de estructura.txt. Los items compuestos (p. ej. "Marca, modelo,
# número de serie") se separan en columnas individuales para facilitar el llenado.
# Se agrega "Host padre" únicamente para vincular VMs con su hipervisor.

IDENTIFICACION = [
    "Hostname", "FQDN", "Alias", "Función o servicios (web, correo, BD, hipervisor...)",
    "Cliente o área dueña", "Responsable técnico", "Responsable de negocio",
    "Criticidad (Tier)", "SLA asociado", "Ventana de mantenimiento", "ID en la CMDB",
]

HARDWARE = [
    "Hostname", "Tipo (físico / VM / cloud / contenedor)", "Host padre",
    "Marca", "Modelo", "Número de serie / Service tag", "CPU", "RAM",
    "Discos (tipo, capacidad, RAID)", "NICs", "Fuentes redundantes",
    "Firmware / BIOS", "iDRAC / iLO / IPMI (IP y versión)",
    "Host / clúster (si es VM)", "Datastore (si es VM)", "Recursos asignados (si es VM)",
    "Garantía y soporte: contrato", "Vencimiento de garantía y soporte", "Proveedor de soporte",
]

UBICACION = [
    "Hostname", "Sitio o data center", "Rack", "Unidad (U)", "Posición",
    "Conexiones de energía (PDU / UPS)",
    "Cableado (patch panel o puerto de switch)",
    "Foto frontal y trasera (referencia)", "Diagrama de rack (referencia)",
]

SISTEMA_OPERATIVO = [
    "Hostname", "Distribución o edición", "Versión", "Kernel o build", "Licenciamiento",
    "Particionado", "Filesystems / puntos de montaje", "LVM", "Zona horaria", "NTP", "Locale",
    "Política de parches", "Fecha del último update",
]

RED = [
    "Hostname", "Interfaz", "IP (v4 / v6)", "Máscara", "Gateway", "VLAN", "DNS", "Dominio",
    "Registros DNS asociados (A, PTR, MX...)", "Bonding / teaming", "MTU", "Rutas estáticas",
    "Puertos en escucha", "Servicios expuestos",
    "Diagrama L2/L3 (switch, puerto, firewall, NAT/VIP) (referencia)",
]

SOFTWARE = [
    "Hostname", "Aplicación instalada", "Versión",
    "Servicio activo / arranque automático",
    "Dependencias (otros servidores, APIs, BD, almacenamiento externo)",
    "Rutas de configuración", "Rutas de logs",
    "Certificados TLS (emisor, CN/SAN, vencimiento, renovación)",
]

SEGURIDAD = [
    "Hostname", "Hardening aplicado (CIS o baseline interno)",
    "Firewall local (reglas)", "Firewall perimetral relacionado",
    "Antimalware / EDR", "IDS", "fail2ban",
    "Usuarios, grupos y sudoers", "Método de autenticación (llaves SSH, MFA, LDAP/AD)",
    "Acceso remoto permitido (IPs, bastión, VPN)",
    "Gestión de contraseñas (bóveda donde se resguardan)",
    "Resultado del último escaneo de vulnerabilidades",
]

RESPALDO = [
    "Hostname", "Qué se respalda", "Herramienta", "Frecuencia", "Retención", "Destino",
    "RPO", "RTO", "Procedimiento de restauración", "Fecha de la última prueba de restauración",
    "Snapshots (si aplica)", "Plan de DR o servidor espejo",
]

MONITOREO = [
    "Hostname", "Herramienta (Zabbix, PRTG, SNMP...)", "Métricas monitoreadas",
    "Umbrales y alertas", "A quién notifican",
    "Envío de logs a SIEM o syslog central",
    "Línea base de desempeño (CPU, RAM, disco, red)",
]

OPERACION = [
    "Hostname", "Procedimiento de arranque", "Procedimiento de apagado ordenado",
    "Tareas programadas (cron / tareas de Windows)",
    "Procedimientos de mantenimiento recurrente", "Runbook de incidentes comunes",
    "Matriz de escalamiento y contactos (interno y proveedores)",
]

CUMPLIMIENTO = [
    "Hostname", "Controles ISO 27001 o normativa aplicable", "Riesgos identificados",
    "Tratamiento de riesgos", "Cambios relevantes (referencias RFC/CHG)",
    "Fin de soporte de hardware", "Fin de soporte del SO", "Reemplazo previsto",
]

ANEXOS = [
    "Hostname", "Tipo de anexo (diagrama, evidencia, configuración, foto, contrato...)",
    "Descripción", "Ubicación / referencia", "Fecha",
]

CONTROL_DOCUMENTO = [
    "Título", "Código", "Versión", "Fecha", "Autor", "Revisor", "Aprobador",
    "Historial de cambios", "Clasificación de la información (confidencial / uso interno)",
]

# Hoja de resumen (Propuesta A): identificación + lo esencial de las demás secciones
RESUMEN = IDENTIFICACION + [
    "Tipo (físico / VM / cloud / contenedor)", "Host padre",
    "Marca", "Modelo", "Número de serie / Service tag", "CPU", "RAM",
    "Discos (tipo, capacidad, RAID)", "NICs",
    "Distribución o edición del SO", "Versión del SO",
    "Sitio o data center", "Rack", "IP principal",
    "Herramienta de respaldo", "Herramienta de monitoreo",
]

# Inventarios de la Propuesta C (físicos y VMs por separado)
FISICO = [
    "Hostname", "FQDN", "Alias", "Función o servicios (web, correo, BD, hipervisor...)",
    "Cliente o área dueña", "Responsable técnico", "Responsable de negocio",
    "Criticidad (Tier)", "SLA asociado", "Ventana de mantenimiento", "ID en la CMDB",
    "Marca", "Modelo", "Número de serie / Service tag", "CPU", "RAM",
    "Discos (tipo, capacidad, RAID)", "NICs", "Firmware / BIOS",
    "iDRAC / iLO / IPMI (IP y versión)",
    "Sitio o data center", "Rack", "Unidad (U)",
    "Garantía y soporte: contrato", "Vencimiento de garantía y soporte", "Proveedor de soporte",
]

VM = [
    "Hostname", "FQDN", "Alias", "Función o servicios (web, correo, BD, hipervisor...)",
    "Cliente o área dueña", "Responsable técnico", "Responsable de negocio",
    "Criticidad (Tier)", "SLA asociado", "Ventana de mantenimiento", "ID en la CMDB",
    "Host padre", "Host / clúster", "Datastore", "Recursos asignados (vCPU, RAM, disco)",
    "CPU", "RAM", "Discos (tipo, capacidad, RAID)", "NICs", "Sitio o data center",
]


# ---------------------------------------------------------------- propuesta A
def generar_propuesta_a(ruta):
    wb = Workbook()
    crear_hoja(wb, "Resumen general", RESUMEN, primera=True)
    crear_hoja(wb, "Hardware y plataforma", HARDWARE)
    crear_hoja(wb, "Ubicación física", UBICACION)
    crear_hoja(wb, "Sistema operativo", SISTEMA_OPERATIVO)
    crear_hoja(wb, "Red (una fila por interfaz)", RED)
    crear_hoja(wb, "Software y servicios (por app)", SOFTWARE)
    crear_hoja(wb, "Seguridad", SEGURIDAD)
    crear_hoja(wb, "Respaldo y recuperación", RESPALDO)
    crear_hoja(wb, "Monitoreo y logging", MONITOREO)
    crear_hoja(wb, "Operación", OPERACION)
    crear_hoja(wb, "Cumplimiento y gestión", CUMPLIMIENTO)
    crear_hoja(wb, "Anexos", ANEXOS)
    crear_hoja(wb, "Control del documento", CONTROL_DOCUMENTO)
    wb.save(ruta)
    return ruta


# ---------------------------------------------------------------- propuesta B
def generar_propuesta_b(ruta):
    wb = Workbook()
    crear_hoja(wb, "1. Control del documento", CONTROL_DOCUMENTO, primera=True)
    crear_hoja(wb, "2. Identificación y propósito", IDENTIFICACION + ["Host padre"])
    crear_hoja(wb, "3. Hardware - plataforma", HARDWARE)
    crear_hoja(wb, "4. Ubicación física", UBICACION)
    crear_hoja(wb, "5. Sistema operativo", SISTEMA_OPERATIVO)
    crear_hoja(wb, "6. Red (una fila por interfaz)", RED)
    crear_hoja(wb, "7. Software y servicios", SOFTWARE)
    crear_hoja(wb, "8. Seguridad", SEGURIDAD)
    crear_hoja(wb, "9. Respaldo y recuperación", RESPALDO)
    crear_hoja(wb, "10. Monitoreo y logging", MONITOREO)
    crear_hoja(wb, "11. Operación", OPERACION)
    crear_hoja(wb, "12. Cumplimiento y gestión", CUMPLIMIENTO)
    crear_hoja(wb, "13. Anexos", ANEXOS)
    wb.save(ruta)
    return ruta


# ---------------------------------------------------------------- propuesta C
def generar_propuesta_c(ruta):
    wb = Workbook()
    crear_hoja(wb, "Inventario (Físicos)", FISICO, primera=True)
    crear_hoja(wb, "Inventario (VMs)", VM)
    crear_hoja(wb, "Red (una fila por interfaz)", RED)
    crear_hoja(wb, "Seguridad", SEGURIDAD)
    crear_hoja(wb, "Respaldo y recuperación", RESPALDO)
    crear_hoja(wb, "Monitoreo y logging", MONITOREO)
    crear_hoja(wb, "Operación", OPERACION)
    crear_hoja(wb, "Anexos", ANEXOS)
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    for ruta in (
        generar_propuesta_a(RUTA_SALIDA_A),
        generar_propuesta_b(RUTA_SALIDA_B),
        generar_propuesta_c(RUTA_SALIDA_C),
    ):
        print(f"Generado: {ruta}")
